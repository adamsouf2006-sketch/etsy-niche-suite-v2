# Render a sheet range to HTML (values from recalculated file, styles from source) for product photos.
import sys, datetime as dt, html, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
SRC = load_workbook('out_example.xlsx'); VAL = load_workbook('rc/out_example.xlsx', data_only=True)
BAR = {'Budget': ['F'], 'Savings & Debt': ['F'], 'Habits': ['AI'], 'Dashboard': ['I'], 'Goals': ['F']}
def color(c):
    try:
        v = c.rgb
        return '#' + v[-6:] if isinstance(v, str) and v not in ('00000000',) else None
    except Exception: return None
def fmt(v, nf):
    if v is None: return ''
    if isinstance(v, (dt.datetime, dt.date)):
        if 'mmmm yyyy' in nf: return v.strftime('%B %Y')
        if 'dddd' in nf: return v.strftime('%A %-d %b %Y')
        if nf == 'd': return str(v.day)
        if 'mmm yy' in nf: return v.strftime('%b %y')
        if nf == 'mmm': return v.strftime('%b')
        if 'd mmm' in nf: return v.strftime('%-d %b')
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if '%' in nf: return f'{v*100:.0f}%'
        if '0.00' in nf and ';' in nf and v == 0: return '–'
        if '0.00' in nf: return f'{v:,.2f}'
        if nf == '0': return f'{v:.0f}'
        return f'{v:g}'
    return str(v)
def colw(s, c):
    for k, d in s.column_dimensions.items():
        if d.min and d.max and d.min <= c <= d.max:
            return 0 if d.hidden else (d.width or 8.43)
    return 8.43
def render(sheet, rng, scale=1.0):
    s, v = SRC[sheet], VAL[sheet]
    a, b = rng.split(':')
    import re
    c1, r1 = re.match(r'([A-Z]+)(\d+)', a).groups(); c2, r2 = re.match(r'([A-Z]+)(\d+)', b).groups()
    c1, c2, r1, r2 = CI(c1), CI(c2), int(r1), int(r2)
    merged = {}; skip = set()
    for m in s.merged_cells.ranges:
        if m.min_row >= r1 and m.min_col >= c1:
            merged[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
            for rr in range(m.min_row, m.max_row + 1):
                for cc in range(m.min_col, m.max_col + 1):
                    if (rr, cc) != (m.min_row, m.min_col): skip.add((rr, cc))
    tw = sum(colw(s, c) * 7.2 * scale for c in range(c1, c2 + 1))
    out = [f'<table class="sh" style="width:{tw:.0f}px"><colgroup>']
    for c in range(c1, c2 + 1):
        w = colw(s, c)
        out.append(f'<col style="width:{w*7.2*scale:.0f}px">')
    out.append('</colgroup>')
    for r in range(r1, r2 + 1):
        h = s.row_dimensions[r].height or 15
        out.append(f'<tr style="height:{h*1.34*scale:.0f}px">')
        for c in range(c1, c2 + 1):
            if (r, c) in skip: continue
            if colw(s, c) == 0: continue
            sc, vc = s.cell(r, c), v.cell(r, c)
            st = []
            bg = color(sc.fill.fgColor) if sc.fill and sc.fill.fill_type == 'solid' else None
            val = vc.value
            if sheet == 'Habits' and val == 'x' and c > 2: bg = '#A8C3A0'
            if bg: st.append(f'background:{bg}')
            f = sc.font
            if f:
                if f.b: st.append('font-weight:700')
                if f.i: st.append('font-style:italic')
                if f.sz: st.append(f'font-size:{f.sz*1.34*scale:.1f}px')
                fc = color(f.color) if f.color else None
                if fc: st.append(f'color:{fc}')
            if sc.border and sc.border.left and sc.border.left.style: st.append('border:1px solid #D9D4CC')
            al = sc.alignment.horizontal if sc.alignment else None
            if al: st.append(f'text-align:{al}')
            if sc.alignment and sc.alignment.wrap_text: st.append('white-space:normal;line-height:1.25')
            elif isinstance(val, (int, float)) and not isinstance(val, bool): st.append('text-align:right')
            txt = html.escape(fmt(val, sc.number_format or ''))
            col = L(c)
            if col in BAR.get(sheet, []) and isinstance(val, (int, float)) and '%' in (sc.number_format or ''):
                p = max(0, min(1, val)) * 100
                barc = '#E07A5F' if (sheet == 'Budget' and r > 20) or (sheet == 'Savings & Debt' and r > 20) else '#A8C3A0'
                st.append(f'background:linear-gradient(90deg,{barc} {p:.0f}%,transparent {p:.0f}%)')
            span = ''
            if (r, c) in merged:
                rs, cs = merged[(r, c)]; span = f' rowspan="{rs}" colspan="{cs}"'
            elif isinstance(val, str) and val and not (sc.alignment and sc.alignment.wrap_text):
                need = len(val) * (f.sz or 10) * 0.8; have = colw(s, c) * 7.2; cs = 1; cc = c + 1
                while have < need and cc <= c2 and (r, cc) not in merged and (r, cc) not in skip and v.cell(r, cc).value in (None, ''):
                    have += colw(s, cc) * 7.2; skip.add((r, cc)); cs += 1; cc += 1
                if cs > 1: span = f' colspan="{cs}"'
            out.append(f'<td{span} style="{";".join(st)}">{txt}</td>')
        out.append('</tr>')
    out.append('</table>')
    return ''.join(out)
CSS = '''.sh{border-collapse:collapse;table-layout:fixed;font-family:Arial,Liberation Sans,sans-serif;color:#1E2B2F;background:#fff}
.sh td{padding:0 6px;overflow:hidden;white-space:nowrap;text-overflow:clip;vertical-align:middle}'''
if __name__ == '__main__':
    sheet, rng, out = sys.argv[1:4]
    open(out, 'w').write(f'<html><head><style>body{{margin:0}}{CSS}</style></head><body>{render(sheet, rng)}</body></html>')
