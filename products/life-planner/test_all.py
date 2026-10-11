import json, re, zipfile, subprocess, datetime as dt, warnings, os, sys
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from PIL import Image
fails = []; ok = 0
def check(cond, msg):
    global ok
    if cond: ok += 1
    else: fails.append(msg)
# 1. recalc both workbooks, no error values
os.makedirs('rc', exist_ok=True)
for f in ['out_example.xlsx', 'out_blank.xlsx']:
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', 'rc', f], capture_output=True)
    wb = load_workbook('rc/' + f, data_only=True); errs = []
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.value.startswith('#') or c.value.startswith('Err:')): errs.append(f'{ws.title}!{c.coordinate}={c.value}')
    check(not errs, f'{f}: error values {errs[:10]}')
    check(len(wb.sheetnames) == 14, f'{f}: {len(wb.sheetnames)} tabs')
# 2. independent maths on example
v = load_workbook('rc/out_example.xlsx', data_only=True)
ml = v['Money Log']; tx = []
for r in range(7, 507):
    d = ml[f'B{r}'].value
    if d: tx.append((d.date() if hasattr(d, 'date') else d, ml[f'C{r}'].value, ml[f'D{r}'].value, ml[f'E{r}'].value))
pm = dt.date(2026, 10, 1); b = v['Budget']
inc = sum(a for d, t, c, a in tx if t == 'Income' and d.month == 10); exp = sum(a for d, t, c, a in tx if t == 'Expense' and d.month == 10)
check(abs(b['D42'].value - inc) < .01 and abs(b['D43'].value - exp) < .01 and abs(b['D44'].value - (inc - exp)) < .01, 'Budget totals')
for r in range(21, 41):
    cat = b[f'B{r}'].value
    if cat:
        e = sum(a for d, t, c, a in tx if t == 'Expense' and c == cat and d.month == 10)
        check(abs((b[f'D{r}'].value or 0) - e) < .01, f'Budget actual {cat}')
        p = b[f'C{r}'].value or 0
        check(abs((b[f'E{r}'].value or 0) - (p - e)) < .01, f'Budget left {cat}')
an = v['Annual']
for r in range(8, 18):
    cat = an[f'B{r}'].value
    if cat:
        for m in range(1, 13):
            e = sum(a for d, t, c, a in tx if t == 'Income' and c == cat and d.month == m)
            val = an.cell(r, 2 + m).value or 0
            check(abs(val - e) < .01, f'Annual {cat} m{m}')
        check(abs((an[f'O{r}'].value or 0) - sum(a for d, t, c, a in tx if t == 'Income' and c == cat)) < .01, f'Annual year {cat}')
sd = v['Savings & Debt']
for r in range(7, 19):
    if sd[f'B{r}'].value:
        check(abs(sd[f'E{r}'].value - max(0, sd[f'C{r}'].value - sd[f'D{r}'].value)) < .01, f'Savings needed r{r}')
        check(abs(sd[f'F{r}'].value - min(1, sd[f'D{r}'].value / sd[f'C{r}'].value)) < 1e-6, f'Savings % r{r}')
check(abs(sd['C36'].value - sum(sd[f'D{r}'].value or 0 for r in range(7, 19))) < .01, 'Total saved')
h = v['Habits']
for r in range(7, 27):
    if h[f'B{r}'].value:
        n = sum(1 for c in range(3, 34) if h.cell(r, c).value == 'x')
        check(h[f'AH{r}'].value == n, f'Habit count r{r}')
        today = dt.date.today(); end = today.day if today.month == 10 else 31; s = 0
        for c in range(2 + end, 2, -1):
            if h.cell(r, c).value == 'x': s += 1
            else: break
        check(h[f'AJ{r}'].value == s, f'Habit streak r{r}: {h[f"AJ{r}"].value} vs {s}')
cl = v['Cleaning']; fq = {'Daily': 1, 'Weekly': 7, 'Fortnightly': 14, 'Monthly': 30, 'Quarterly': 91, 'Half-yearly': 182, 'Yearly': 365}
for r in range(7, 47):
    if cl[f'B{r}'].value and cl[f'E{r}'].value:
        nd = cl[f'E{r}'].value + dt.timedelta(fq[cl[f'C{r}'].value])
        check(cl[f'F{r}'].value == nd, f'Cleaning next r{r}')
tk = v['Tasks']; open_t = sorted([(tk[f'D{r}'].value, tk[f'B{r}'].value) for r in range(7, 107) if tk[f'B{r}'].value and tk[f'F{r}'].value != 'Done'])
ds = v['Dashboard']
check([ds[f'E{r}'].value for r in range(10, 10 + min(8, len(open_t)))] == [t for _, t in open_t[:8]], 'Dashboard next tasks order')
top = sorted([(b[f'D{r}'].value, b[f'B{r}'].value) for r in range(21, 41) if b[f'B{r}'].value and b[f'D{r}'].value], key=lambda x: -x[0])[:5]
check([ds[f'C{r}'].value for r in range(14, 19)] == [a for a, _ in top], 'Dashboard top spending values')
check(len(set(ds[f'B{r}'].value for r in range(14, 19))) == 5, 'Dashboard top spending unique')
cal = v['Calendar']; days = [cal.cell(7 + w * 2, c).value for w in range(6) for c in range(2, 9)]
days = [d for d in days if d]
check(len(days) == 31 and days[0].day == 1 and days[0].weekday() == 3, 'Calendar Oct 2026 starts Thursday, 31 days')
# 3. blank workbook empty of example data
bl = load_workbook('rc/out_blank.xlsx', data_only=True)
check(all(bl['Money Log'][f'B{r}'].value is None for r in range(7, 507)), 'Blank: money log empty')
check(all(bl['Tasks'][f'B{r}'].value is None for r in range(7, 107)), 'Blank: tasks empty')
check(bl['Budget']['D42'].value == 0 and bl['Dashboard']['E10'].value in (None, ''), 'Blank: dashboard clean')
# 4. formulas use only Sheets+Excel-compatible functions
ALLOWED = {'IF', 'SUM', 'SUMIFS', 'COUNTIF', 'COUNTIFS', 'IFERROR', 'INDEX', 'MATCH', 'SMALL', 'LARGE', 'INT', 'ROW', 'COLUMN', 'OFFSET', 'MAX', 'MIN', 'N', 'OR', 'AND',
           'DATE', 'TEXT', 'ABS', 'YEAR', 'MONTH', 'DAY', 'EOMONTH', 'TODAY', 'WEEKDAY'}
src = load_workbook('out_example.xlsx'); used = set(); bad_cf = []
for ws in src:
    for row in ws.iter_rows():
        for c in row:
            f = c.value.text if hasattr(c.value, 'text') else c.value
            if isinstance(f, str) and f.startswith('='): used |= set(re.findall(r'([A-Z][A-Z0-9\.]+)\(', f))
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            for f in (rule.formula or []):
                if '!' in f: bad_cf.append(ws.title)
check(used <= ALLOWED, f'Non-whitelisted functions: {used - ALLOWED}')
check(not bad_cf, f'Cross-sheet conditional formatting: {bad_cf}')
# 5. data validations present
dvs = {ws.title: len(ws.data_validations.dataValidation) for ws in src}
check(dvs['Money Log'] >= 2 and dvs['Tasks'] >= 3 and dvs['Cleaning'] >= 1, f'Data validations {dvs}')
# 6. no "###": formatted width vs column width for value cells
def colw(ws, c):
    for k, d in ws.column_dimensions.items():
        if d.min and d.max and d.min <= c <= d.max: return d.width or 8.43
    return 8.43
merged = {ws.title: {(r, c) for m in ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1)} for ws in src}
narrow = []
for ws in v:
    sws = src[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) and (c.row, c.column) not in merged[ws.title]:
                nf = sws.cell(c.row, c.column).number_format; sz = (sws.cell(c.row, c.column).font.sz or 10)
                txt = f'{c.value*100:.0f}%' if '%' in nf else (f'{c.value:,.2f}' if '0.00' in nf else f'{c.value:g}')
                if len(txt) * sz / 10 * 1.1 + 1 > colw(sws, c.column): narrow.append(f'{ws.title}!{c.coordinate}')
check(not narrow, f'Possible ### cells: {narrow[:10]}')
# 7. photos
for f in sorted(os.listdir('photos')):
    if f.endswith('.jpg') and (f.startswith('p') or f == 'main_c.jpg'):
        im = Image.open('photos/' + f)
        check(im.size == (2000, 2500) and os.path.getsize('photos/' + f) < 20e6, f'Photo {f} {im.size}')
# 8. video
pr = json.loads(subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration:stream=codec_name,codec_type,width,height,profile,pix_fmt', '-of', 'json', 'video/life-planner.mp4'], capture_output=True, text=True).stdout)
vs = [s for s in pr['streams'] if s['codec_type'] == 'video'][0]; au = [s for s in pr['streams'] if s['codec_type'] == 'audio']
check(3 <= float(pr['format']['duration']) <= 15 and vs['width'] == 1080 and vs['codec_name'] == 'h264' and vs['pix_fmt'] == 'yuv420p' and au, f'Video {pr}')
check(os.path.getsize('video/life-planner.mp4') < 100e6, 'Video size')
# 9. listing text
L = json.load(open('listing.json'))
t = L['title']
check(len(t) <= 140 and t.count('&') <= 1 and t.count('%') <= 1 and t.count(':') <= 1, f'Title rules ({len(t)})')
check(sum(1 for w in t.split() if w.isupper() and len(w) > 1) <= 3, 'Title caps words')
check(len(L['tags']) == 13 and all(len(x) <= 20 for x in L['tags']) and len(set(L['tags'])) == 13, f'Tags {[x for x in L["tags"] if len(x)>20]}')
check(not any(x in ' '.join(L['tags']) for x in ['instant download', 'digital download', 'editable']), 'Generic tags')
check(re.fullmatch(r"[A-Za-z0-9 ,&'\-]+", t) is not None and all(re.fullmatch(r"[a-z0-9 '&\-]+", x) for x in L['tags']), 'Allowed characters')
# 10. zip
z = zipfile.ZipFile('Oakfold-Life-Planner.zip'); check(z.testzip() is None and len(z.namelist()) == 3 and os.path.getsize('Oakfold-Life-Planner.zip') < 20e6, 'ZIP integrity')
check(not any(n.startswith('.') or '~' in n for n in z.namelist()), 'ZIP temp files')
# 11. spelling of visible text (US): no British forms
txt = ' '.join(str(c.value) for ws in src for row in ws.iter_rows() for c in row if isinstance(c.value, str) and not c.value.startswith('='))
txt += open('guide.html').read() + open('photos.py').read()
brit = [w for w in ['colour', 'cancelled', 'favourite', 'organise', 'centre', 'Wholemeal', 'Mince', 'Coloured', 'Fortnightly', 'Tinned', 'Bin bags'] if w in txt]
check(not brit, f'British spellings: {brit}')
check('  ' not in re.sub(r'\s{2,}', lambda m: m.group() if '\n' in m.group() else '  ', ' '.join(str(c.value) for ws in src for row in ws.iter_rows() for c in row if isinstance(c.value, str))) , 'Double spaces')
print(f'PASSED {ok}  FAILED {len(fails)}')
for f in fails: print(' -', f)
