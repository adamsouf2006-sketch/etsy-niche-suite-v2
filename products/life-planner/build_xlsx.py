# Oakfold Life Planner Spreadsheet: one workbook for Google Sheets and Excel.
import datetime as dt, random, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

EXAMPLE = '--blank' not in sys.argv
OUT = sys.argv[1]
INK, TEAL, SAND, CREAM, CORAL, SAGE, GREY = '1E2B2F', '1F4E5A', 'E9DCC9', 'FBF7F1', 'E07A5F', 'A8C3A0', 'D9D4CC'
F = 'Arial'
fill = lambda c: PatternFill('solid', fgColor=c)
thin = Side(style='thin', color=GREY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = 'FFFFFF'
random.seed(7)

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, sub, widths, tab=TEAL):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    for r in range(1, 4):
        for c in range(1, len(widths) + 1):
            ws.cell(r, c).fill = fill(TEAL)
    ws['B2'] = title; ws['B2'].font = Font(F, 20, True, color='FFFFFF')
    ws['B3'] = sub; ws['B3'].font = Font(F, 10, color=SAND)
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[4].height = 8
    return ws

def head(ws, row, col, labels, color=INK):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = Font(F, 10, True, color='FFFFFF'); c.fill = fill(color)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = box
    ws.row_dimensions[row].height = 28

def cell(ws, ref, v=None, bold=False, size=10, color=INK, bg=None, fmt=None, align=None, border=True, italic=False):
    c = ws[ref]
    if v is not None: c.value = v
    c.font = Font(F, size, bold, italic, color=color)
    if bg: c.fill = fill(bg)
    if fmt: c.number_format = fmt
    if align: c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    if border: c.border = box
    return c

def grid(ws, r1, r2, c1, c2, bg=INPUT, fmt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c); x.border = box; x.font = Font(F, 10, color=INK)
            if bg: x.fill = fill(bg)
            if fmt: x.number_format = fmt

def dv(ws, formula, rng):
    d = DataValidation(type='list', formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rng)

CUR = '[$€-x-euro2] #,##0.00'  # replaced by setup-driven symbol in labels; keep neutral format below
MONEY = '#,##0.00'
MONEY0 = '#,##0.00;-#,##0.00;"–"'
DATE = 'yyyy-mm-dd'
TODAY = dt.date(2026, 10, 12)

# ---------- Start Here ----------
ws = sheet('Start Here', 'Life Planner', 'Budget • Tasks • Habits • Goals • Meals • Cleaning — Google Sheets & Excel', [3, 30, 70, 3])
rows = [
    ('1. Setup', 'Type your name, currency symbol, the month you are planning and your own budget categories. Two minutes, once.'),
    ('2. Money Log', 'The only place you type money. One row per income or expense. Budget, Annual, Savings and Dashboard update themselves.'),
    ('3. Budget', 'Type your planned amount per category. Actual, difference and "safe to spend per day" are calculated for the month chosen in Setup.'),
    ('4. Annual', 'All 12 months side by side, built from the Money Log. Nothing to type.'),
    ('5. Savings & Debt', 'Up to 12 savings goals and 12 debts with live progress bars.'),
    ('6. Tasks', '100 tasks with category, due date, priority and status. Open tasks appear on the Dashboard, soonest first.'),
    ('7. Habits', '20 habits x 31 days. Type x on the days you did it; streak, total and month score are automatic.'),
    ('8. Cleaning', 'Daily to yearly chores. Type the date you last did it; the next due date and status are calculated.'),
    ('9. Goals', '10 goals with why, reward, five dated steps and a progress bar.'),
    ('10. Week Plan', 'Plan your week by the hour, from 6:00 to 22:00.'),
    ('11. Meals', 'Plan breakfast, lunch, dinner and snacks for the week, with a grocery list by aisle.'),
    ('12. Calendar', 'Choose a month: every day shows how many open tasks are due.'),
    ('White cells', 'are for typing. Colored cells hold formulas: leave them, and the planner never breaks.'),
    ('Google Sheets', 'File > Import > Upload this .xlsx, or open it from Drive and choose File > Save as Google Sheets.'),
]
for i, (a, b) in enumerate(rows):
    r = 6 + i
    cell(ws, f'B{r}', a, True, 11, TEAL, CREAM, align='left')
    cell(ws, f'C{r}', b, size=10, bg='FFFFFF', align='left')
    ws.row_dimensions[r].height = 34
cell(ws, 'B22', 'Oakfold', True, 9, CORAL, border=False)
cell(ws, 'C22', 'Personal use. Not financial advice.', size=9, color='7A7A7A', border=False, italic=True)

# ---------- Setup ----------
ws = sheet('Setup', 'Setup', 'Fill the white cells once. Everything else follows.', [3, 26, 22, 4, 26, 22, 3])
for r, (k, v) in enumerate([('Your name', 'Alex' if EXAMPLE else ''), ('Currency symbol', '$'),
                            ('Planning month (first day)', dt.date(2026, 10, 1)), ('Year', '=YEAR(C8)')], 6):
    cell(ws, f'B{r}', k, True, bg=CREAM, align='left')
    c = cell(ws, f'C{r}', v, bg='FFFFFF' if r < 9 else SAND, align='center')
ws['C8'].number_format = 'mmmm yyyy'
head(ws, 11, 2, ['Income categories'], TEAL); head(ws, 11, 5, ['Expense categories'], CORAL)
inc = ['Salary', 'Side income', 'Gifts', 'Refunds', 'Other income']
exp = ['Rent / Mortgage', 'Utilities', 'Groceries', 'Transport', 'Insurance', 'Phone & Internet', 'Subscriptions',
       'Eating out', 'Health', 'Personal care', 'Clothing', 'Kids', 'Pets', 'Gifts & Events', 'Entertainment',
       'Savings transfer', 'Debt payment', 'Other']
for i in range(10):
    cell(ws, f'B{12+i}', inc[i] if i < len(inc) else None, bg='FFFFFF')
for i in range(20):
    cell(ws, f'E{12+i}', exp[i] if i < len(exp) else None, bg='FFFFFF')
cell(ws, 'B34', 'Tip: rename any category; every tab picks up the new name.', italic=True, color='7A7A7A', border=False)
from openpyxl.workbook.defined_name import DefinedName
for n, ref in [('IncomeCats', "'Setup'!$B$12:$B$21"), ('ExpenseCats', "'Setup'!$E$12:$E$31"),
               ('PlanMonth', "'Setup'!$C$8"), ('Cur', "'Setup'!$C$7")]:
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

# ---------- Money Log ----------
ws = sheet('Money Log', 'Money Log', 'One row per transaction. Choose Income or Expense, then a category.', [3, 13, 11, 22, 13, 34, 11, 3])
head(ws, 6, 2, ['Date', 'Type', 'Category', 'Amount', 'Note', 'Month'])
N = 500
grid(ws, 7, 6 + N, 2, 6); grid(ws, 7, 6 + N, 7, 7, bg=CREAM)
for r in range(7, 7 + N):
    ws[f'B{r}'].number_format = DATE; ws[f'E{r}'].number_format = MONEY
    ws[f'G{r}'] = f'=IF(B{r}="","",DATE(YEAR(B{r}),MONTH(B{r}),1))'; ws[f'G{r}'].number_format = 'mmm yy'
dv(ws, '"Income,Expense"', f'C7:C{6+N}')
dv(ws, "='Setup'!$B$12:$E$31", f'D7:D{6+N}')
if EXAMPLE:
    tx = []
    for m in (8, 9, 10):
        tx += [(dt.date(2026, m, 1), 'Income', 'Salary', 3250), (dt.date(2026, m, 15), 'Income', 'Side income', random.choice([280, 340, 410])),
               (dt.date(2026, m, 1), 'Expense', 'Rent / Mortgage', 1150), (dt.date(2026, m, 5), 'Expense', 'Utilities', random.choice([96, 118, 132])),
               (dt.date(2026, m, 6), 'Expense', 'Phone & Internet', 65), (dt.date(2026, m, 9), 'Expense', 'Insurance', 142),
               (dt.date(2026, m, 12), 'Expense', 'Subscriptions', 38), (dt.date(2026, m, 20), 'Expense', 'Savings transfer', 300),
               (dt.date(2026, m, 21), 'Expense', 'Debt payment', 250)]
        for d in (3, 10, 17, 24):
            if m == 10 and d > 11: continue
            tx += [(dt.date(2026, m, d), 'Expense', 'Groceries', random.choice([84, 96, 112, 127])),
                   (dt.date(2026, m, d + 1), 'Expense', 'Transport', random.choice([35, 42, 48]))]
        tx += [(dt.date(2026, m, 8), 'Expense', 'Eating out', 46), (dt.date(2026, m, 11), 'Expense', 'Entertainment', 29),
               (dt.date(2026, m, 4), 'Expense', 'Health', 24), (dt.date(2026, m, 7), 'Expense', 'Pets', 58)]
    tx = [t for t in tx if t[0] <= TODAY]; tx.sort()
    notes = {'Groceries': 'Weekly shop', 'Salary': 'Payday', 'Rent / Mortgage': 'Rent', 'Debt payment': 'Card payment', 'Savings transfer': 'Emergency fund'}
    for i, (d, t, c, a) in enumerate(tx):
        r = 7 + i
        ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'E{r}'], ws[f'F{r}'] = d, t, c, a, notes.get(c, '')
ws.freeze_panes = 'B7'
ML = f"'Money Log'!$G$7:$G${6+N}"; MT = f"'Money Log'!$C$7:$C${6+N}"; MC = f"'Money Log'!$D$7:$D${6+N}"; MA = f"'Money Log'!$E$7:$E${6+N}"

# ---------- Budget ----------
ws = sheet('Budget', 'Monthly Budget', 'Month shown = Planning month in Setup. Type planned amounts in the white column.', [3, 24, 14, 14, 14, 12, 3])
cell(ws, 'B6', 'Month', True, bg=CREAM); c = cell(ws, 'C6', '=PlanMonth', True, bg=SAND, fmt='mmmm yyyy', align='center')
ws.merge_cells('C6:D6')
head(ws, 8, 2, ['Income', 'Planned', 'Actual', 'Difference', 'Received'], TEAL)
pin = [3250, 350, 0, 0, 0]
for i in range(10):
    r = 9 + i
    ws[f'B{r}'] = f"=IF('Setup'!B{12+i}=\"\",\"\",'Setup'!B{12+i})"
    ws[f'C{r}'] = (pin[i] if i < 5 and pin[i] else None) if EXAMPLE else None
    ws[f'D{r}'] = f'=IF(B{r}="","",SUMIFS({MA},{MC},B{r},{MT},"Income",{ML},PlanMonth))'
    ws[f'E{r}'] = f'=IF(B{r}="","",D{r}-N(C{r}))'
    ws[f'F{r}'] = f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}/C{r})'
grid(ws, 9, 18, 2, 6, bg=CREAM); grid(ws, 9, 18, 3, 3)
for r in range(9, 19):
    for col in 'CDE': ws[f'{col}{r}'].number_format = MONEY
    ws[f'F{r}'].number_format = '0%'
head(ws, 20, 2, ['Expenses', 'Planned', 'Actual', 'Left', 'Used'], CORAL)
pex = [1150, 120, 450, 180, 142, 65, 40, 80, 40, 30, 40, 0, 60, 50, 40, 300, 250, 50]
for i in range(20):
    r = 21 + i
    ws[f'B{r}'] = f"=IF('Setup'!E{12+i}=\"\",\"\",'Setup'!E{12+i})"
    ws[f'C{r}'] = (pex[i] if i < len(pex) and pex[i] else None) if EXAMPLE else None
    ws[f'D{r}'] = f'=IF(B{r}="","",SUMIFS({MA},{MC},B{r},{MT},"Expense",{ML},PlanMonth))'
    ws[f'E{r}'] = f'=IF(B{r}="","",N(C{r})-D{r})'
    ws[f'F{r}'] = f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}/C{r})'
    ws[f'H{r}'] = f'=IF(OR(B{r}="",N(D{r})=0),"",D{r}+ROW()/1000000)'; ws[f'H{r}'].font = Font(F, 8, color='FFFFFF')
grid(ws, 21, 40, 2, 6, bg=CREAM); grid(ws, 21, 40, 3, 3)
for r in range(21, 41):
    for col in 'CDE': ws[f'{col}{r}'].number_format = MONEY
    ws[f'F{r}'].number_format = '0%'
ws.conditional_formatting.add('E21:E40', CellIsRule(operator='lessThan', formula=['0'], font=Font(F, 10, True, color='C0392B')))
ws.conditional_formatting.add('F21:F40', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=CORAL))
for r, (k, f) in enumerate([('Total income', '=SUM(D9:D18)'), ('Total expenses', '=SUM(D21:D40)'), ('Balance', '=D42-D43'),
                            ('Planned left to spend', '=SUM(C21:C40)-D43'), ('Days left this month', '=IF(MONTH(TODAY())=MONTH(PlanMonth),DAY(EOMONTH(PlanMonth,0))-DAY(TODAY())+1,"")'),
                            ('Safe to spend per day', '=IF(N(D46)=0,"",MAX(0,D45)/D46)')], 42):
    cell(ws, f'B{r}', k, True, bg=CREAM); ws.merge_cells(f'B{r}:C{r}')
    cell(ws, f'D{r}', f, True, bg=SAND, fmt=MONEY if r != 46 else '0', align='right')

# ---------- Annual ----------
ws = sheet('Annual', 'Annual Overview', 'Twelve months from the Money Log for the year in Setup. Nothing to type.', [3, 22] + [10.5] * 12 + [12, 3])
head(ws, 6, 2, ['Category'] + [dt.date(2026, m, 1).strftime('%b') for m in range(1, 13)] + ['Year'])
for m in range(12):
    ws.cell(5, 3 + m, f"=DATE('Setup'!$C$9,{m+1},1)").number_format = 'mmm'
    ws.cell(5, 3 + m).font = Font(F, 8, color='FFFFFF')
r = 7
for label, typ, src, n in [('INCOME', 'Income', 'B', 10), ('EXPENSES', 'Expense', 'E', 20)]:
    cell(ws, f'B{r}', label, True, color='FFFFFF', bg=TEAL if typ == 'Income' else CORAL); r += 1
    first = r
    for i in range(n):
        ws[f'B{r}'] = f"=IF('Setup'!{src}{12+i}=\"\",\"\",'Setup'!{src}{12+i})"
        for m in range(12):
            col = L(3 + m)
            ws[f'{col}{r}'] = f'=IF($B{r}="","",SUMIFS({MA},{MC},$B{r},{MT},"{typ}",{ML},{col}$5))'
        ws[f'O{r}'] = f'=IF($B{r}="","",SUM(C{r}:N{r}))'
        r += 1
    grid(ws, first, r - 1, 2, 15, bg=CREAM, fmt=MONEY0)
    cell(ws, f'B{r}', f'Total {typ.lower()}', True, bg=SAND)
    for m in range(13):
        col = L(3 + m); cell(ws, f'{col}{r}', f'=SUM({col}{first}:{col}{r-1})', True, bg=SAND, fmt=MONEY0)
    if typ == 'Income': TI = r
    else: TE = r
    r += 2
cell(ws, f'B{r}', 'Net (income − expenses)', True, color='FFFFFF', bg=INK)
for m in range(13):
    col = L(3 + m); cell(ws, f'{col}{r}', f'={col}{TI}-{col}{TE}', True, bg=SAND, fmt=MONEY0)
ws.freeze_panes = 'C7'

# ---------- Savings & Debt ----------
ws = sheet('Savings & Debt', 'Savings & Debt', 'Type targets, balances and amounts saved or paid. Progress fills itself.', [3, 24, 13, 13, 13, 12, 13, 3])
head(ws, 6, 2, ['Savings goal', 'Target', 'Saved', 'Still needed', 'Progress', 'Target date'], TEAL)
sg = [('Emergency fund', 5000, 2100, dt.date(2027, 6, 30)), ('Vacation', 1800, 650, dt.date(2027, 7, 1)), ('New laptop', 1200, 900, dt.date(2027, 1, 31)), ('Car repairs', 800, 240, None)]
for i in range(12):
    r = 7 + i
    if EXAMPLE and i < len(sg): ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'G{r}'] = sg[i]
    ws[f'E{r}'] = f'=IF(B{r}="","",MAX(0,N(C{r})-N(D{r})))'; ws[f'F{r}'] = f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,N(D{r})/C{r}))'
grid(ws, 7, 18, 2, 7); grid(ws, 7, 18, 5, 6, bg=CREAM)
head(ws, 21, 2, ['Debt', 'Starting balance', 'Paid so far', 'Remaining', 'Paid off', 'Min. payment'], CORAL)
db = [('Credit card', 3200, 1150, 90), ('Car loan', 9800, 2600, 245), ('Student loan', 14500, 1800, 160)]
for i in range(12):
    r = 22 + i
    if EXAMPLE and i < len(db): ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'G{r}'] = db[i]
    ws[f'E{r}'] = f'=IF(B{r}="","",MAX(0,N(C{r})-N(D{r})))'; ws[f'F{r}'] = f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,N(D{r})/C{r}))'
grid(ws, 22, 33, 2, 7); grid(ws, 22, 33, 5, 6, bg=CREAM)
for r in list(range(7, 19)) + list(range(22, 34)):
    for col in 'CDE': ws[f'{col}{r}'].number_format = MONEY
    ws[f'F{r}'].number_format = '0%'
    ws[f'G{r}'].number_format = DATE if r < 19 else MONEY
ws.conditional_formatting.add('F7:F18', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=SAGE))
ws.conditional_formatting.add('F22:F33', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=CORAL))
for r, (k, f) in enumerate([('Total saved', '=SUM(D7:D18)'), ('Total debt remaining', '=SUM(E22:E33)'), ('Monthly minimum payments', '=SUM(G22:G33)')], 36):
    cell(ws, f'B{r}', k, True, bg=CREAM); cell(ws, f'C{r}', f, True, bg=SAND, fmt=MONEY)

# ---------- Tasks ----------
ws = sheet('Tasks', 'Task List', 'Type a task, pick a category, a due date and a priority. Mark it Done when finished.', [3, 36, 14, 13, 11, 12, 10, 3])
head(ws, 6, 2, ['Task', 'Category', 'Due date', 'Priority', 'Status', 'Days left'])
NT = 100
grid(ws, 7, 6 + NT, 2, 6); grid(ws, 7, 6 + NT, 7, 7, bg=CREAM)
for r in range(7, 7 + NT):
    ws[f'D{r}'].number_format = DATE
    ws[f'G{r}'] = f'=IF(OR(B{r}="",D{r}="",F{r}="Done"),"",D{r}-TODAY())'
    ws[f'I{r}'] = f'=IF(OR(B{r}="",D{r}="",F{r}="Done"),"",D{r}+ROW()/100000)'
    ws[f'I{r}'].font = Font(F, 8, color='FFFFFF')
ws.column_dimensions['I'].hidden = True
dv(ws, '"Home,Work,Money,Health,Family,Errands,Personal"', f'C7:C{6+NT}')
dv(ws, '"High,Medium,Low"', f'E7:E{6+NT}'); dv(ws, '"To do,In progress,Done"', f'F7:F{6+NT}')
ws.conditional_formatting.add(f'B7:G{6+NT}', FormulaRule(formula=['$F7="Done"'], font=Font(F, 10, color='9A9A9A', strike=True)))
ws.conditional_formatting.add(f'G7:G{6+NT}', FormulaRule(formula=['AND($G7<>"",$G7<0)'], font=Font(F, 10, True, color='C0392B')))
ws.conditional_formatting.add(f'E7:E{6+NT}', CellIsRule(operator='equal', formula=['"High"'], fill=fill('F6D5CC')))
if EXAMPLE:
    tk = [('Pay car insurance', 'Money', 2, 'High', 'To do'), ('Book dentist check-up', 'Health', 5, 'Medium', 'To do'),
          ('Renew passport', 'Personal', 18, 'Medium', 'In progress'), ('Order school photos', 'Family', 3, 'Low', 'To do'),
          ('Send quarterly report', 'Work', 1, 'High', 'In progress'), ('Return online order', 'Errands', -1, 'Medium', 'To do'),
          ('Clear out garage', 'Home', 12, 'Low', 'To do'), ('Call the bank about fees', 'Money', 4, 'Medium', 'To do'),
          ('Plan weekend trip', 'Family', 9, 'Low', 'To do'), ('Change smoke alarm batteries', 'Home', -3, 'High', 'Done')]
    for i, (t, c, d, p, s) in enumerate(tk):
        r = 7 + i; ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'E{r}'], ws[f'F{r}'] = t, c, TODAY + dt.timedelta(d), p, s
ws.freeze_panes = 'B7'

# ---------- Habits ----------
ws = sheet('Habits', 'Habit Tracker', 'Type x on each day you did the habit. Month shown = Planning month in Setup.', [3, 24] + [3.4] * 31 + [7, 7, 7, 3])
head(ws, 6, 2, ['Habit'])
for d in range(31):
    c = ws.cell(6, 3 + d, f'=IF({d+1}>DAY(EOMONTH(PlanMonth,0)),"",{d+1})')
    c.font = Font(F, 8, True, color='FFFFFF'); c.fill = fill(INK); c.alignment = Alignment(horizontal='center'); c.border = box
head(ws, 6, 34, ['Done', 'Score', 'Streak'])
hb = ['Drink 2 L of water', 'Walk 8,000 steps', 'Read 20 minutes', 'No phone after 22:00', 'Stretch / yoga', 'Make the bed',
      'Cook at home', 'Journal', 'Take vitamins', 'Tidy for 10 minutes', 'No spend day', 'Sleep by 23:00']
NH = 20
for i in range(NH):
    r = 7 + i
    if EXAMPLE and i < len(hb):
        ws[f'B{r}'] = hb[i]
        for d in range(11):
            if random.random() < 0.72: ws.cell(r, 3 + d, 'x')
    ws[f'AH{r}'] = f'=IF(B{r}="","",COUNTIF(C{r}:AG{r},"x"))'
    ws[f'AI{r}'] = f'=IF(B{r}="","",AH{r}/MAX(1,IF(MONTH(TODAY())=MONTH(PlanMonth),DAY(TODAY()),DAY(EOMONTH(PlanMonth,0)))))'
    # streak: consecutive x ending today (or month end)
    endc = f'IF(MONTH(TODAY())=MONTH(PlanMonth),DAY(TODAY()),DAY(EOMONTH(PlanMonth,0)))'
    ws[f'AJ{r}'] = ArrayFormula(f'AJ{r}', f'=IF(B{r}="","",IFERROR({endc}-MAX(IF(OFFSET(C{r},0,0,1,{endc})<>"x",COLUMN(OFFSET(C{r},0,0,1,{endc}))-COLUMN(C{r})+1,0)),0))')
grid(ws, 7, 6 + NH, 2, 33); grid(ws, 7, 6 + NH, 34, 36, bg=CREAM)
for r in range(7, 7 + NH):
    for c in range(3, 34): ws.cell(r, c).alignment = Alignment(horizontal='center')
    ws[f'AI{r}'].number_format = '0%'
ws.conditional_formatting.add(f'C7:AG{6+NH}', CellIsRule(operator='equal', formula=['"x"'], fill=fill(SAGE), font=Font(F, 10, True, color=TEAL)))
ws.conditional_formatting.add(f'AI7:AI{6+NH}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=SAGE))
ws.freeze_panes = 'C7'

# ---------- Cleaning ----------
ws = sheet('Cleaning', 'Cleaning Schedule', 'Type the date you last did each chore. Next due date and status update themselves.', [3, 32, 14, 13, 13, 13, 12, 3])
head(ws, 6, 2, ['Chore', 'Frequency', 'Room', 'Last done', 'Next due', 'Status'])
cl = [('Wipe kitchen counters', 'Daily', 'Kitchen', 0), ('Load / empty dishwasher', 'Daily', 'Kitchen', 0), ('Quick bathroom wipe', 'Daily', 'Bathroom', -1),
      ('Vacuum floors', 'Weekly', 'Whole home', -4), ('Change bed sheets', 'Weekly', 'Bedroom', -8), ('Clean toilets', 'Weekly', 'Bathroom', -6),
      ('Mop floors', 'Weekly', 'Whole home', -5), ('Clean out fridge', 'Weekly', 'Kitchen', -9), ('Dust surfaces', 'Weekly', 'Living room', -3),
      ('Clean oven', 'Monthly', 'Kitchen', -40), ('Wash windows inside', 'Monthly', 'Whole home', -20), ('Descale kettle', 'Monthly', 'Kitchen', -15),
      ('Wash duvets & pillows', 'Quarterly', 'Bedroom', -70), ('Clean behind appliances', 'Quarterly', 'Kitchen', -95),
      ('Deep clean carpets', 'Half-yearly', 'Living room', -120), ('Clear gutters', 'Yearly', 'Outside', -300)]
NC = 40
dv(ws, '"Daily,Weekly,Every 2 weeks,Monthly,Quarterly,Half-yearly,Yearly"', f'C7:C{6+NC}')
for i in range(NC):
    r = 7 + i
    if i < len(cl):
        ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'] = cl[i][:3]
        if EXAMPLE: ws[f'E{r}'] = TODAY + dt.timedelta(cl[i][3])
    ws[f'E{r}'].number_format = DATE
    days = f'IF(C{r}="Daily",1,IF(C{r}="Weekly",7,IF(C{r}="Every 2 weeks",14,IF(C{r}="Monthly",30,IF(C{r}="Quarterly",91,IF(C{r}="Half-yearly",182,365))))))'
    ws[f'F{r}'] = f'=IF(OR(B{r}="",C{r}=""),"",IF(E{r}="",TODAY(),E{r}+{days}))'; ws[f'F{r}'].number_format = DATE
    ws[f'G{r}'] = f'=IF(F{r}="","",IF(F{r}<TODAY(),"Overdue",IF(F{r}=TODAY(),"Today","Upcoming")))'
grid(ws, 7, 6 + NC, 2, 5); grid(ws, 7, 6 + NC, 6, 7, bg=CREAM)
for r in range(7, 7 + NC):
    ws[f'E{r}'].number_format = DATE; ws[f'F{r}'].number_format = DATE
ws.conditional_formatting.add(f'G7:G{6+NC}', CellIsRule(operator='equal', formula=['"Overdue"'], fill=fill('F6D5CC'), font=Font(F, 10, True, color='C0392B')))
ws.conditional_formatting.add(f'G7:G{6+NC}', CellIsRule(operator='equal', formula=['"Today"'], fill=fill('FCEBC4')))

# ---------- Goals ----------
ws = sheet('Goals', 'Goal Tracker', 'Ten goals. For each: why it matters, five dated steps, and a reward. Tick steps with x.', [3, 26, 30, 13, 6, 14, 3])
NG = 10
r = 6
ex = [('Run a 10K', 'Feel fit and proud', ['Buy running shoes', 'Run 3K without stopping', 'Run 5K', 'Run 8K', 'Race day'], 'New running watch'),
      ('Save $5,000 emergency fund', 'Stop worrying about surprise bills', ['Open savings account', 'Automate $300 a month', 'Reach $2,500', 'Reach $4,000', 'Reach $5,000'], 'Weekend away')]
gstart = []
for g in range(NG):
    head(ws, r, 2, ['Goal', 'Why it matters', 'Deadline', 'Done', 'Progress'], TEAL if g % 2 == 0 else CORAL)
    gstart.append(r)
    grid(ws, r + 1, r + 1, 2, 4); grid(ws, r + 1, r + 1, 5, 6, bg=CREAM)
    ws[f'D{r+1}'].number_format = DATE
    ws[f'F{r+1}'] = f'=IF(B{r+1}="","",COUNTIF(E{r+2}:E{r+6},"x")/5)'; ws[f'F{r+1}'].number_format = '0%'
    for s in range(5):
        rr = r + 2 + s
        cell(ws, f'B{rr}', f'Step {s+1}', italic=True, color='7A7A7A', bg=CREAM)
        grid(ws, rr, rr, 3, 5); ws[f'D{rr}'].number_format = DATE; ws[f'E{rr}'].alignment = Alignment(horizontal='center')
    cell(ws, f'B{r+7}', 'Reward', italic=True, color='7A7A7A', bg=CREAM); grid(ws, r + 7, r + 7, 3, 3)
    if EXAMPLE and g < len(ex):
        n, w, st, rw = ex[g]
        ws[f'B{r+1}'], ws[f'C{r+1}'], ws[f'D{r+1}'] = n, w, dt.date(2027, 5 + g, 30)
        for s in range(5):
            ws[f'C{r+2+s}'] = st[s]; ws[f'D{r+2+s}'] = TODAY + dt.timedelta(30 * (s + 1) - 40)
            if s < 2 - g: ws[f'E{r+2+s}'] = 'x'
        ws[f'C{r+7}'] = rw
    ws.conditional_formatting.add(f'F{r+1}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=SAGE))
    r += 9

# ---------- Week Plan ----------
days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
ws = sheet('Week Plan', 'Weekly Planner', 'Plan your week by the hour. Type the Monday date once.', [3, 9] + [17] * 7 + [3])
cell(ws, 'B5', 'Week of', True, bg=CREAM); ws.merge_cells('C5:D5')
cell(ws, 'C5', '=TODAY()-WEEKDAY(TODAY(),3)', True, bg='FFFFFF', fmt='dddd d mmm yyyy', align='center')
head(ws, 6, 2, ['Time'] + days)
for d in range(7):
    c = ws.cell(7, 3 + d, f'=$C$5+{d}'); c.number_format = 'd mmm'; c.font = Font(F, 9, color=TEAL); c.alignment = Alignment(horizontal='center'); c.fill = fill(CREAM); c.border = box
cell(ws, 'B7', '', bg=CREAM)
for h in range(17):
    r = 8 + h
    cell(ws, f'B{r}', f'{6+h:02d}:00', True, 9, TEAL, CREAM, align='center')
    grid(ws, r, r, 3, 9)
    ws.row_dimensions[r].height = 22
if EXAMPLE:
    for (d, h, t) in [(0, 7, 'Gym'), (0, 9, 'Team meeting'), (1, 12, 'Lunch with Sam'), (2, 18, 'Meal prep'), (3, 9, 'Report due'),
                      (4, 17, 'Groceries'), (5, 10, 'Kids swim class'), (6, 9, 'Long run'), (6, 16, 'Plan next week')]:
        ws.cell(8 + h - 6, 3 + d, t)
for i, t in enumerate(['Top 3 this week', 'Notes']):
    cell(ws, f'B{27+i*5}', t, True, color='FFFFFF', bg=CORAL); ws.merge_cells(f'B{27+i*5}:I{27+i*5}')
    grid(ws, 28 + i * 5, 30 + i * 5, 2, 9)

# ---------- Meals ----------
ws = sheet('Meals', 'Meal Planner & Grocery List', 'Plan the week, then list what to buy by aisle. Tick items with x as you shop.', [3, 12] + [17] * 7 + [3])
head(ws, 6, 2, [''] + days)
meals = ['Breakfast', 'Lunch', 'Dinner', 'Snacks']
ex_m = {'Breakfast': ['Oats & berries', 'Yogurt & granola', 'Eggs on toast', 'Oats & berries', 'Smoothie', 'Pancakes', 'Eggs & avocado'],
        'Lunch': ['Chicken wrap', 'Pasta leftovers', 'Tuna salad', 'Soup & bread', 'Grain bowl', 'Picnic', 'Roast leftovers'],
        'Dinner': ['Veggie stir-fry', 'Spaghetti bolognese', 'Fish tacos', 'Chicken curry', 'Homemade pizza', 'Burgers', 'Sunday roast'],
        'Snacks': ['Apple & nuts', 'Hummus & carrots', 'Banana', 'Cheese & crackers', 'Popcorn', 'Fruit salad', 'Dark chocolate']}
for i, m in enumerate(meals):
    r = 7 + i
    cell(ws, f'B{r}', m, True, color=TEAL, bg=CREAM, align='center'); grid(ws, r, r, 3, 9)
    ws.row_dimensions[r].height = 34
    for d in range(7):
        ws.cell(r, 3 + d).alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        if EXAMPLE: ws.cell(r, 3 + d, ex_m[m][d])
aisles = ['Produce', 'Meat & Fish', 'Dairy & Eggs', 'Bakery', 'Pantry', 'Frozen', 'Household']
ex_g = {'Produce': ['Berries', 'Carrots', 'Avocados', 'Onions', 'Spinach'], 'Meat & Fish': ['Chicken thighs', 'Ground beef', 'White fish'],
        'Dairy & Eggs': ['Eggs (12)', 'Greek yogurt', 'Cheddar'], 'Bakery': ['Whole wheat bread', 'Tortilla wraps'],
        'Pantry': ['Oats', 'Spaghetti', 'Curry paste', 'Canned tomatoes'], 'Frozen': ['Peas', 'Pizza bases'], 'Household': ['Dish soap', 'Trash bags']}
cell(ws, 'B13', 'Grocery list', True, 12, 'FFFFFF', CORAL); ws.merge_cells('B13:I13')
for j, a in enumerate(aisles):
    col = 3 + j
    c = ws.cell(14, col, a); c.font = Font(F, 10, True, color='FFFFFF'); c.fill = fill(TEAL); c.alignment = Alignment(horizontal='center'); c.border = box
    for k in range(15):
        x = ws.cell(15 + k, col); x.border = box; x.font = Font(F, 10, color=INK)
        if EXAMPLE and k < len(ex_g[a]): x.value = ex_g[a][k]
cell(ws, 'B14', 'Aisle', True, color='FFFFFF', bg=INK)
for k in range(15): cell(ws, f'B{15+k}', '', bg=CREAM)

# ---------- Calendar ----------
ws = sheet('Calendar', 'Monthly Calendar', 'Month = Planning month in Setup. Each day counts tasks due (Tasks tab).', [3] + [16] * 7 + [3])
cell(ws, 'B5', '=PlanMonth', True, 14, TEAL, border=False, fmt='mmmm yyyy'); ws.merge_cells('B5:D5')
head(ws, 6, 2, ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun'])
for w in range(6):
    for d in range(7):
        r = 7 + w * 2; col = L(2 + d); k = w * 7 + d
        f = f'=IF(MONTH(PlanMonth-WEEKDAY(PlanMonth,3)+{k})<>MONTH(PlanMonth),"",PlanMonth-WEEKDAY(PlanMonth,3)+{k})'
        c = cell(ws, f'{col}{r}', f, True, 11, TEAL, CREAM, 'd', 'left')
        t = f"=IF({col}{r}=\"\",\"\",IF(COUNTIFS('Tasks'!$D$7:$D$106,{col}{r},'Tasks'!$F$7:$F$106,\"<>Done\")=0,\"\",COUNTIFS('Tasks'!$D$7:$D$106,{col}{r},'Tasks'!$F$7:$F$106,\"<>Done\")&\" due\"))"
        cell(ws, f'{col}{r+1}', t, size=9, color=CORAL, align='left')
        ws.row_dimensions[r + 1].height = 38
    ws.conditional_formatting.add(f'B{7+w*2}:H{7+w*2}', FormulaRule(formula=[f'B{7+w*2}=TODAY()'], fill=fill(SAND)))

# ---------- Dashboard (first) ----------
ws = sheet('Dashboard', 'Dashboard', '=IF(\'Setup\'!C6="","Your month at a glance","Hi "&\'Setup\'!C6&" — your month at a glance")', [3, 22, 15, 3, 30, 12, 3, 22, 12, 3])
def kpi(ref, label, f, color):
    c1, r = ref[0], int(ref[1:])
    cell(ws, f'{c1}{r}', label, True, 9, 'FFFFFF', color, align='center')
    c = cell(ws, f'{c1}{r+1}', f, True, 18, color, 'FFFFFF', MONEY, 'center'); ws.row_dimensions[r + 1].height = 34
cell(ws, 'B5', '=PlanMonth', True, 12, TEAL, border=False, fmt='mmmm yyyy')
kpi('B6', 'Income', "=Cur&TEXT('Budget'!D42,\"#,##0.00\")", TEAL); kpi('E6', 'Expenses', "=Cur&TEXT('Budget'!D43,\"#,##0.00\")", CORAL); kpi('H6', 'Balance', "=IF('Budget'!D44<0,\"-\","")&Cur&TEXT(ABS('Budget'!D44),\"#,##0.00\")", INK)
ws.merge_cells('B6:C6'); ws.merge_cells('B7:C7'); ws.merge_cells('E6:F6'); ws.merge_cells('E7:F7'); ws.merge_cells('H6:I6'); ws.merge_cells('H7:I7')
cell(ws, 'B9', 'Safe to spend / day', True, bg=CREAM); cell(ws, 'C9', "=IF('Budget'!D47=\"\",\"\",Cur&TEXT('Budget'!D47,\"#,##0.00\"))", True, bg=SAND, align='right')
cell(ws, 'B10', 'Total saved', True, bg=CREAM); cell(ws, 'C10', "=Cur&TEXT('Savings & Debt'!C36,\"#,##0.00\")", True, bg=SAND, align='right')
cell(ws, 'B11', 'Debt remaining', True, bg=CREAM); cell(ws, 'C11', "=Cur&TEXT('Savings & Debt'!C37,\"#,##0.00\")", True, bg=SAND, align='right')
head(ws, 13, 2, ['Top spending', 'Actual'], CORAL); ws.row_dimensions[13].height = 20
for k in range(5):
    r = 14 + k
    ws[f'C{r}'] = f"=IFERROR(INDEX('Budget'!$D$21:$D$40,MATCH(LARGE('Budget'!$H$21:$H$40,{k+1}),'Budget'!$H$21:$H$40,0)),\"\")"
    ws[f'B{r}'] = f"=IFERROR(INDEX('Budget'!$B$21:$B$40,MATCH(LARGE('Budget'!$H$21:$H$40,{k+1}),'Budget'!$H$21:$H$40,0)),\"\")"
grid(ws, 14, 18, 2, 3, bg=CREAM, fmt=MONEY)
head(ws, 9, 5, ['Next tasks', 'Due'], TEAL)
for k in range(8):
    r = 10 + k
    ws[f'F{r}'] = f"=IFERROR(INT(SMALL('Tasks'!$I$7:$I$106,{k+1})),\"\")"
    ws[f'E{r}'] = f"=IFERROR(INDEX('Tasks'!$B$7:$B$106,MATCH(SMALL('Tasks'!$I$7:$I$106,{k+1}),'Tasks'!$I$7:$I$106,0)),\"\")"
grid(ws, 10, 17, 5, 6, bg=CREAM); 
for r in range(10, 18): ws[f'F{r}'].number_format = 'd mmm'
ws.conditional_formatting.add('F10:F17', FormulaRule(formula=['AND(F10<>"",F10<TODAY())'], font=Font(F, 10, True, color='C0392B')))
head(ws, 9, 8, ['Habits this month', 'Score'], INK)
for k in range(8):
    r = 10 + k
    ws[f'H{r}'] = f"=IF('Habits'!B{7+k}=\"\",\"\",'Habits'!B{7+k})"; ws[f'I{r}'] = f"=IF('Habits'!B{7+k}=\"\",\"\",'Habits'!AI{7+k})"
grid(ws, 10, 17, 8, 9, bg=CREAM)
for r in range(10, 18): ws[f'I{r}'].number_format = '0%'
ws.conditional_formatting.add('I10:I17', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=SAGE))
head(ws, 20, 5, ['Cleaning due', 'Status'], CORAL)
for k in range(6):
    r = 21 + k
    ws[f'F{r}'] = f"=IFERROR(INDEX('Cleaning'!$G$7:$G$46,SMALL(IF('Cleaning'!$G$7:$G$46<>\"Upcoming\",IF('Cleaning'!$G$7:$G$46<>\"\",ROW('Cleaning'!$G$7:$G$46)-6)),{k+1})),\"\")"
    ws[f'E{r}'] = f"=IFERROR(INDEX('Cleaning'!$B$7:$B$46,SMALL(IF('Cleaning'!$G$7:$G$46<>\"Upcoming\",IF('Cleaning'!$G$7:$G$46<>\"\",ROW('Cleaning'!$G$7:$G$46)-6)),{k+1})),\"\")"
grid(ws, 21, 26, 5, 6, bg=CREAM)
head(ws, 20, 8, ['Savings goals', 'Progress'], TEAL)
for k in range(6):
    r = 21 + k
    ws[f'H{r}'] = f"=IF('Savings & Debt'!B{7+k}=\"\",\"\",'Savings & Debt'!B{7+k})"; ws[f'I{r}'] = f"=IF('Savings & Debt'!B{7+k}=\"\",\"\",'Savings & Debt'!F{7+k})"
grid(ws, 21, 26, 8, 9, bg=CREAM)
for r in range(21, 27): ws[f'I{r}'].number_format = '0%'
ws.conditional_formatting.add('I21:I26', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=SAGE))
head(ws, 20, 2, ['Meals today', ''], INK)
for k, m in enumerate(meals):
    r = 21 + k
    cell(ws, f'B{r}', m, bg=CREAM)
    cell(ws, f'C{r}', f"=INDEX('Meals'!$C${7+k}:$I${7+k},WEEKDAY(TODAY(),2))", bg=CREAM, align='left')
wb.move_sheet('Dashboard', offset=-(len(wb.sheetnames) - 2))

# array formulas for cleaning-due list
from openpyxl.worksheet.formula import ArrayFormula
for k in range(6):
    r = 21 + k
    for col in 'EF':
        f = ws[f'{col}{r}'].value
        ws[f'{col}{r}'] = ArrayFormula(f'{col}{r}', f)

for s in wb.worksheets:
    s.sheet_view.zoomScale = 100
    s.page_setup.orientation = 'landscape'; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
wb['Budget'].column_dimensions['H'].hidden = True
wb.active = 1
wb.save(OUT)
print('saved', OUT, wb.sheetnames)
