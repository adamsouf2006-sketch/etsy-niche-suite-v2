# HTML -> PDF (Chromium) + champs remplissables (PyMuPDF) + Word (python-docx), en Letter et A4.
import sys,os,asyncio,html as H; sys.path.insert(0,os.path.abspath("../../commun"))
from content import *
import pymupdf
from docx import Document
from docx.shared import Pt, RGBColor, Mm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SIZES={"US-Letter":(8.5,11),"A4":(8.27,11.69)}
CSS=f'''@page{{margin:0}}*{{box-sizing:border-box;margin:0}}body{{font-family:Carlito,Calibri,sans-serif;color:#1F2937;font-size:10.5pt}}
.pg{{width:var(--w);height:var(--h);overflow:hidden;padding:.55in .6in .5in;position:relative;page-break-after:always}}
.top{{display:flex;align-items:flex-end;justify-content:space-between;border-bottom:3px solid {ACC};padding-bottom:8px;margin-bottom:10px}}
.top h1{{font-size:24pt;color:{ACC};line-height:1}} .top small{{color:#6B7280;font-size:10pt}} .logo{{border:1.5px dashed #9CA3AF;color:#9CA3AF;font-size:8pt;width:1.3in;height:.55in;display:grid;place-items:center;border-radius:6px}}
h2{{font-size:11.5pt;color:#fff;background:{ACC};padding:3px 8px;border-radius:4px;margin:calc(10px*var(--s)) 0 calc(6px*var(--s));letter-spacing:.03em;text-transform:uppercase}}
.row{{display:flex;gap:10px;margin:calc(5px*var(--s)) 0}} .fld{{display:flex;flex-direction:column}} .fld label{{font-size:8pt;color:#6B7280;font-weight:700;text-transform:uppercase;letter-spacing:.04em}}
.box{{height:calc(20px*var(--s));border-bottom:1.2px solid #9CA3AF;background:{SOFT}33}} .big .box{{height:calc(34px*var(--s))}}
.cks{{display:grid;gap:calc(5px*var(--s)) 14px;margin:calc(4px*var(--s)) 0}} .ck{{display:flex;align-items:center;gap:7px;font-size:10pt}} .cb{{width:12px;height:12px;border:1.3px solid {ACC};border-radius:2px;flex:none}}
table{{width:100%;border-collapse:collapse;margin:4px 0;font-size:9.5pt}} th{{background:{SOFT};color:{ACC};text-align:left;padding:4px 6px;font-size:8.5pt;text-transform:uppercase}} td{{border-bottom:1px solid #D1D5DB;height:calc(20px*var(--s));padding:0}}
ol{{margin:2px 0 4px 16px;font-size:9.3pt;line-height:1.3}} li{{margin:calc(2.5px*var(--s)) 0}}
.sigs{{display:flex;gap:26px;margin-top:calc(16px*var(--s))}} .sig{{flex:1}} .sig .line{{height:30px;border-bottom:1.2px solid #374151}} .sig small{{font-size:8pt;color:#6B7280}}
.note{{font-size:8pt;color:#6B7280;margin-top:8px;font-style:italic}}
.foot{{position:absolute;bottom:.28in;left:.6in;right:.6in;display:flex;justify-content:space-between;font-size:7.5pt;color:#9CA3AF}}'''
def fid(): fid.n+=1; return f"f{fid.n}"
fid.n=0
def html_doc(num,title,sub,blocks):
    o=[f'<div class=pg><div class=top><div><small>{H.escape(sub)}</small><h1>{H.escape(title)}</h1></div><div class=logo>YOUR LOGO</div></div>']
    for b in blocks:
        k=b[0]
        if k=="h": o.append(f"<h2>{H.escape(b[1])}</h2>")
        elif k=="f":
            o.append('<div class=row>'+"".join(f'<div class="fld{" big" if "notes" in l.lower() or "details" in l.lower() and w==3 else ""}" style="flex:{w}"><label>{H.escape(l)}</label><div class=box data-f="{fid()}" data-l="{H.escape(l)}"></div></div>' for l,w in b[1])+'</div>')
        elif k=="c": o.append(f'<div class=cks style="grid-template-columns:repeat({b[1]},1fr)">'+"".join(f'<div class=ck><div class=cb data-c="{fid()}" data-l="{H.escape(x)}"></div>{H.escape(x)}</div>' for x in b[2])+'</div>')
        elif k=="t":
            cols=b[1]; o.append("<table><tr>"+"".join(f"<th>{H.escape(c)}</th>" for c in cols)+"</tr>"+"".join("<tr>"+"".join(f'<td data-f="{fid()}" data-l="{H.escape(c)} {r+1}"></td>' for c in cols)+"</tr>" for r in range(b[2]))+"</table>")
        elif k=="clauses": o.append("<ol>"+"".join(f"<li>{H.escape(x)}</li>" for x in b[1])+"</ol>")
        elif k=="sig": o.append('<div class=sigs>'+"".join(f'<div class=sig><div class=line data-f="{fid()}" data-l="{H.escape(s)} signature"></div><small>{H.escape(s)} - signature</small></div><div class=sig style="flex:.45"><div class=line data-f="{fid()}" data-l="{H.escape(s)} date"></div><small>Date</small></div>' for s in b[1])+'</div>')
        elif k=="note": o.append(f'<div class=note>{H.escape(b[1])}</div>')
    o.append('<div class=end></div>'); o.append(f'<div class=foot><span>{BRAND} · {KIT}</span><span>{num} / {len(DOCS):02d}</span></div></div>')
    return "".join(o)
async def build(size):
    from playwright.async_api import async_playwright
    w,h=SIZES[size]; os.makedirs(f"../out/{size}",exist_ok=True); report=[]
    async with async_playwright() as p:
        br=await p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
        for num,title,sub,blocks in DOCS:
            fid.n=0; body=html_doc(num,title,sub,blocks)
            page=await br.new_page(viewport={"width":int(w*96),"height":int(h*96)})
            await page.set_content(f"<html><head><meta charset=utf-8><style>{CSS}</style></head><body style='--w:{w}in;--h:{h}in'>{body}</body></html>")
            await page.wait_for_timeout(150)
            sc=1.0
            for _ in range(6):
                await page.evaluate(f"document.body.style.setProperty('--s','{sc}')")
                end=await page.evaluate("document.querySelector('.end').getBoundingClientRect().top")
                avail=h*96-0.75*96
                if abs(avail-end)<12: break
                sc=max(1.0,min(2.2,sc*(1+(avail-end)/end*1.6)))
            hgt=end; await page.evaluate(f"document.body.style.setProperty('--s','{sc}')")
            rects=await page.evaluate("""[...document.querySelectorAll('[data-f],[data-c]')].map(e=>{const r=e.getBoundingClientRect();return {f:e.dataset.f||e.dataset.c,c:!!e.dataset.c,l:e.dataset.l,x:r.x,y:r.y,w:r.width,h:r.height}})""")
            fn=f"../out/{size}/{num} - {title}"
            await page.pdf(path=fn+"_raw.pdf",width=f"{w}in",height=f"{h}in",print_background=True)
            await page.close()
            doc=pymupdf.open(fn+"_raw.pdf"); pg=doc[0]; s=72/96; n=0
            for r in rects:
                if r["y"]*s>pg.rect.height: continue
                wd=pymupdf.Widget(); rect=pymupdf.Rect(r["x"]*s,r["y"]*s,(r["x"]+r["w"])*s,(r["y"]+r["h"])*s)
                if r["c"]: wd.field_type=pymupdf.PDF_WIDGET_TYPE_CHECKBOX; wd.field_value=False
                else: wd.field_type=pymupdf.PDF_WIDGET_TYPE_TEXT; wd.text_fontsize=0 if r["h"]>25 else 9; wd.text_font="Helv"
                if not r["c"] and r["h"]>25: wd.field_flags=pymupdf.PDF_TX_FIELD_IS_MULTILINE
                wd.field_name=f'{r["l"]} [{r["f"]}]'; wd.rect=rect; wd.border_width=0; pg.add_widget(wd); n+=1
            doc.save(fn+".pdf"); os.remove(fn+"_raw.pdf")
            report.append((size,num,title,len(doc),n,len(rects),round(sc,2),round(hgt/96,2)))
        await br.close()
    return report
def docx_doc(size,num,title,sub,blocks):
    w,h=SIZES[size]; d=Document(); sec=d.sections[0]; sec.page_width=Inches(w); sec.page_height=Inches(h)
    for m in ("left_margin","right_margin"): setattr(sec,m,Inches(.6))
    sec.top_margin=Inches(.5); sec.bottom_margin=Inches(.45)
    st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10); st.paragraph_format.space_after=Pt(2)
    acc=RGBColor.from_string(ACC[1:])
    def shade(cell,hexc):
        tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),hexc); tcPr.append(sh)
    p=d.add_paragraph(); r=p.add_run(sub); r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x6B,0x72,0x80)
    p=d.add_paragraph(); r=p.add_run(title); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=acc
    p.add_run("\t[YOUR LOGO]").font.size=Pt(8)
    for b in blocks:
        k=b[0]
        if k=="h":
            t=d.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,ACC[1:]); rr=c.paragraphs[0].add_run(b[1].upper()); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255); rr.font.size=Pt(10.5)
        elif k=="f":
            t=d.add_table(rows=2,cols=len(b[1]))
            for i,(l,wd) in enumerate(b[1]):
                rr=t.cell(0,i).paragraphs[0].add_run(l.upper()); rr.font.size=Pt(7.5); rr.bold=True; rr.font.color.rgb=RGBColor(0x6B,0x72,0x80)
                shade(t.cell(1,i),"EEF7F5"); t.cell(1,i).paragraphs[0].add_run(" ")
        elif k=="c":
            n=b[1]; items=b[2]; rows=(len(items)+n-1)//n; t=d.add_table(rows=rows,cols=n)
            for i,x in enumerate(items): t.cell(i//n,i%n).paragraphs[0].add_run("☐  "+x).font.size=Pt(10)
        elif k=="t":
            t=d.add_table(rows=b[2]+1,cols=len(b[1])); t.style="Table Grid"
            for i,c in enumerate(b[1]): shade(t.cell(0,i),"E6F4F1"); rr=t.cell(0,i).paragraphs[0].add_run(c.upper()); rr.bold=True; rr.font.size=Pt(8); rr.font.color.rgb=acc
        elif k=="clauses":
            for i,x in enumerate(b[1],1): q=d.add_paragraph(f"{i}. {x}"); q.paragraph_format.left_indent=Inches(.15); q.runs[0].font.size=Pt(9.3)
        elif k=="sig":
            for s in b[1]: d.add_paragraph(f"{s} signature: ______________________________    Date: ____________")
        elif k=="note": q=d.add_paragraph(b[1]); q.runs[0].italic=True; q.runs[0].font.size=Pt(8)
        if k in ("f","t","c"): sp=d.add_paragraph(); sp.paragraph_format.space_after=Pt(0); sp.runs or None; sp.paragraph_format.line_spacing=Pt(4)
    ft=sec.footer.paragraphs[0]; ft.text=f"{BRAND} · {KIT} · {num}/{len(DOCS):02d}"; ft.runs[0].font.size=Pt(7.5)
    d.save(f"../out/{size}/{num} - {title}.docx")
if __name__=="__main__":
    rep=[]
    for size in SIZES:
        rep+=asyncio.run(build(size))
        for num,t,s,b in DOCS: docx_doc(size,num,t,s,b)
    for r in rep: print(r)
