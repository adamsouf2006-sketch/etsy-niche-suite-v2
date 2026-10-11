import datetime as dt, data
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, DoughnutChart, Reference
N="1E3A5F"; NL="E6EEF7"; GO="C9A227"; INK="1B2430"; MUT="667085"
F=lambda **k: Font(name="Arial",color=k.pop("color",INK),**k); fill=lambda c: PatternFill("solid",fgColor=c)
B=Border(bottom=Side(style="thin",color="DCE3EC"))
MON=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]; MONEY='$#,##0.00;[Red]-$#,##0.00'
wb=Workbook(); wb.remove(wb.active)
def sheet(name,widths,title,sub):
    ws=wb.create_sheet(name); ws.sheet_properties.tabColor=N; ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=3
    for i,w in enumerate(widths): ws.column_dimensions[chr(66+i)].width=w
    ws["B2"]=title; ws["B2"].font=F(size=20,bold=True,color=N); ws["B3"]=sub; ws["B3"].font=F(size=10,color=MUT); ws.row_dimensions[2].height=30
    return ws
def head(ws,row,cols,c0=2):
    for i,h in enumerate(cols):
        c=ws.cell(row,c0+i,h); c.font=F(bold=True,color="FFFFFF",size=10); c.fill=fill(N); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[row].height=24
# Start
ws=sheet("Start Here",[95],"Small Business Bookkeeping "+str(data.YEAR),"Read me first - 3 minutes")
txt=["1. SETUP: type your business name, tax set-aside %, mileage rate and your income / expense categories (blue cells).",
"2. INCOME: log every payment you receive. Category and payment method are drop-downs. The month fills in by itself.",
"3. EXPENSES: log every business purchase with its category. Mark it deductible (Y/N).",
"4. MILEAGE: log business trips. The deduction is calculated with your rate.",
"5. PROFIT & LOSS: your monthly and yearly profit and loss statement, by category - automatic.",
"6. DASHBOARD: income, expenses, net profit, margin, tax to set aside, charts - automatic.",
"7. QUARTERLY TAX: estimated tax to set aside each quarter, with typical US due dates.","",
"Only type in blue-shaded cells. White cells contain formulas.",
"Works in Google Sheets (upload to Drive, then Open with Google Sheets) and Excel 2016 / Microsoft 365.",
"Sample data is included so you can see it working. Select the blue cells and press Delete when you are ready.","",
"This is a general bookkeeping tool, not tax or legal advice. Check rates and deadlines with your accountant."]
for i,s in enumerate(txt): c=ws.cell(5+i,2,s); c.font=F(size=11,bold=s[:1].isdigit()); c.alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[5+i].height=30 if s else 10
# Setup
ws=sheet("Setup",[30,18,4,30],"Setup","Your business details and categories")
for r,(l,v,fmt) in enumerate([("Business name",data.BIZ,None),("Year",data.YEAR,"0"),("Tax set-aside %",data.TAX,"0%"),("Mileage rate ($/mile)",data.MILE,'$0.000')],5):
    ws.cell(r,2,l).font=F(bold=True); c=ws.cell(r,3,v); c.fill=fill(NL)
    if fmt: c.number_format=fmt
head(ws,10,["Income categories"]); head(ws,10,["Expense categories"],5)
for i in range(10):
    c=ws.cell(11+i,2,data.INC_CATS[i] if i<len(data.INC_CATS) else None); c.fill=fill(NL)
for i in range(25):
    c=ws.cell(11+i,5,data.EXP_CATS[i] if i<len(data.EXP_CATS) else None); c.fill=fill(NL)
INC_R="Setup!$B$11:$B$20"; EXP_R="Setup!$E$11:$E$35"
# Income
TMAX=1500
ws=sheet("Income",[13,22,16,18,14,16,8],"Income","Every payment you receive")
head(ws,5,["Date","Client","Category","Reference","Amount","Payment method","Month"])
dvc=DataValidation(type="list",formula1="="+INC_R,allow_blank=True); dvp=DataValidation(type="list",formula1='"Bank transfer,Card,Check,Cash,PayPal,Other"',allow_blank=True)
ws.add_data_validation(dvc); ws.add_data_validation(dvp); r=6
for m in range(1,10):
    for d,cl,cat,ref,a,pm in data.income(m):
        for col,v in zip(range(2,8),(dt.date(data.YEAR,m,d),cl,cat,ref,a,pm)): ws.cell(r,col,v)
        r+=1
for rr in range(6,TMAX+1):
    for col in range(2,8): ws.cell(rr,col).fill=fill(NL)
    ws.cell(rr,2).number_format="mmm d, yyyy"; ws.cell(rr,6).number_format=MONEY
    ws.cell(rr,8,f'=IF(B{rr}="","",TEXT(B{rr},"mmm"))').alignment=Alignment(horizontal="center")
dvc.add(f"D6:D{TMAX}"); dvp.add(f"G6:G{TMAX}"); ws.freeze_panes="B6"
# Expenses
ws=sheet("Expenses",[13,22,20,20,14,12,8],"Expenses","Every business purchase")
head(ws,5,["Date","Vendor","Category","Note","Amount","Deductible","Month"])
dve=DataValidation(type="list",formula1="="+EXP_R,allow_blank=True); dvy=DataValidation(type="list",formula1='"Y,N"',allow_blank=True)
ws.add_data_validation(dve); ws.add_data_validation(dvy); r=6
for m in range(1,10):
    for d,v,cat,note,a in data.expenses(m):
        for col,val in zip(range(2,8),(dt.date(data.YEAR,m,d),v,cat,note,a,"Y")): ws.cell(r,col,val)
        r+=1
for rr in range(6,TMAX+1):
    for col in range(2,8): ws.cell(rr,col).fill=fill(NL)
    ws.cell(rr,2).number_format="mmm d, yyyy"; ws.cell(rr,6).number_format=MONEY; ws.cell(rr,7).alignment=Alignment(horizontal="center")
    ws.cell(rr,8,f'=IF(B{rr}="","",TEXT(B{rr},"mmm"))').alignment=Alignment(horizontal="center")
dve.add(f"D6:D{TMAX}"); dvy.add(f"G6:G{TMAX}"); ws.freeze_panes="B6"
# Mileage
ws=sheet("Mileage",[13,20,22,12,14,8],"Mileage Log","Business trips and deduction")
head(ws,5,["Date","Purpose","Client / place","Miles","Deduction","Month"]); r=6
for m in range(1,10):
    for d,p,c,mi in data.mileage(m):
        ws.cell(r,2,dt.date(data.YEAR,m,d)); ws.cell(r,3,p); ws.cell(r,4,c); ws.cell(r,5,mi); r+=1
for rr in range(6,600):
    for col in range(2,6): ws.cell(rr,col).fill=fill(NL)
    ws.cell(rr,2).number_format="mmm d, yyyy"; ws.cell(rr,5).number_format="0.0"
    ws.cell(rr,6,f'=IF(E{rr}="","",E{rr}*Setup!$C$8)').number_format=MONEY
    ws.cell(rr,7,f'=IF(B{rr}="","",TEXT(B{rr},"mmm"))').alignment=Alignment(horizontal="center")
ws["I5"]="Total miles"; ws["I6"]="=SUM(E6:E599)"; ws["J5"]="Total deduction"; ws["J6"]="=SUM(F6:F599)"; ws["J6"].number_format=MONEY
for c in ("I5","J5"): ws[c].font=F(bold=True,color=N)
ws.column_dimensions["I"].width=14; ws.column_dimensions["J"].width=16
I=lambda col: f"Income!${col}$6:${col}${TMAX}"; E=lambda col: f"Expenses!${col}$6:${col}${TMAX}"
# P&L
ws=sheet("Profit & Loss",[24]+[11]*12+[13],"Profit & Loss Statement","Monthly and year to date")
head(ws,5,["Category"]+MON+["Year total"])
r=6; ws.cell(r,2,"INCOME").font=F(bold=True,color=N); r+=1; i0=r
for i in range(10):
    ws.cell(r,2,f'=IF(Setup!B{11+i}="","",Setup!B{11+i})')
    for m in range(12):
        L=ws.cell(5,3+m).column_letter
        ws.cell(r,3+m,f'=IF($B{r}="","",SUMIFS({I("F")},{I("D")},$B{r},{I("H")},{L}$5))').number_format='#,##0;-#,##0;"-"'
    ws.cell(r,15,f'=IF(B{r}="","",SUM(C{r}:N{r}))').number_format='$#,##0'; r+=1
i1=r-1; ws.cell(r,2,"Total income").font=F(bold=True); TI=r
for m in range(13): L=ws.cell(5,3+m).column_letter; c=ws.cell(r,3+m,f"=SUM({L}{i0}:{L}{i1})"); c.number_format='$#,##0'; c.font=F(bold=True); c.fill=fill(NL)
r+=2; ws.cell(r,2,"EXPENSES").font=F(bold=True,color=N); r+=1; e0=r
for i in range(25):
    ws.cell(r,2,f'=IF(Setup!E{11+i}="","",Setup!E{11+i})')
    for m in range(12):
        L=ws.cell(5,3+m).column_letter
        ws.cell(r,3+m,f'=IF($B{r}="","",SUMIFS({E("F")},{E("D")},$B{r},{E("H")},{L}$5))').number_format='#,##0;-#,##0;"-"'
    ws.cell(r,15,f'=IF(B{r}="","",SUM(C{r}:N{r}))').number_format='$#,##0'; r+=1
e1=r-1; ws.cell(r,2,"Mileage deduction")
for m in range(12): L=ws.cell(5,3+m).column_letter; ws.cell(r,3+m,f'=SUMIFS(Mileage!$F$6:$F$599,Mileage!$G$6:$G$599,{L}$5)').number_format='#,##0;-#,##0;"-"'
ws.cell(r,15,f"=SUM(C{r}:N{r})").number_format='$#,##0'; MI=r; r+=1
ws.cell(r,2,"Total expenses").font=F(bold=True); TE=r
for m in range(13): L=ws.cell(5,3+m).column_letter; c=ws.cell(r,3+m,f"=SUM({L}{e0}:{L}{MI})"); c.number_format='$#,##0'; c.font=F(bold=True); c.fill=fill(NL)
r+=2; ws.cell(r,2,"NET PROFIT").font=F(bold=True,color=N,size=12); NP=r
for m in range(13): L=ws.cell(5,3+m).column_letter; c=ws.cell(r,3+m,f"={L}{TI}-{L}{TE}"); c.number_format='$#,##0;[Red]-$#,##0'; c.font=F(bold=True,color=N); c.fill=fill("FFF4D6")
r+=1; ws.cell(r,2,"Profit margin"); PM=r
for m in range(13): L=ws.cell(5,3+m).column_letter; ws.cell(r,3+m,f'=IFERROR({L}{NP}/{L}{TI},"")').number_format="0%"
ws.freeze_panes="C6"
bc=BarChart(); bc.title="Income vs expenses"; bc.add_data(Reference(ws,min_col=2,max_col=14,min_row=TI),from_rows=True,titles_from_data=True); bc.add_data(Reference(ws,min_col=2,max_col=14,min_row=TE),from_rows=True,titles_from_data=True)
bc.set_categories(Reference(ws,min_col=3,max_col=14,min_row=5)); bc.height=8; bc.width=24; ws.add_chart(bc,f"B{r+3}")
PL="'Profit & Loss'!"
# Dashboard
ws=sheet("Dashboard",[22,16,16,16,16,4,18],"Business Dashboard","Year to date - updates automatically")
ws["B5"]=f'=Setup!C5'; ws["B5"].font=F(bold=True,size=13,color=N)
labels=[("INCOME",f"={PL}O{TI}"),("EXPENSES",f"={PL}O{TE}"),("NET PROFIT",f"={PL}O{NP}"),("TAX TO SET ASIDE",f"=MAX(0,{PL}O{NP})*Setup!C7")]
for i,(l,f) in enumerate(labels):
    c=ws.cell(7,2+i,l); c.font=F(size=9,bold=True,color=MUT); c.alignment=Alignment(horizontal="center")
    c=ws.cell(8,2+i,f); c.number_format='$#,##0;[Red]-$#,##0'; c.font=F(size=17,bold=True,color=N); c.alignment=Alignment(horizontal="center"); c.fill=fill(NL)
ws.row_dimensions[8].height=36
ws["B10"]="Profit margin"; ws["C10"]=f"={PL}O{PM}"; ws["C10"].number_format="0.0%"; ws["C10"].font=F(bold=True,color=N)
ws["B11"]="Miles driven"; ws["C11"]="=Mileage!I6"; ws["C11"].number_format="#,##0"
head(ws,13,["Month","Income","Expenses","Net profit"])
for m in range(12):
    r=14+m; L=chr(67+m) if m<12 else "O"
    from openpyxl.utils import get_column_letter as gl
    L=gl(3+m); ws.cell(r,2,MON[m])
    for j,row in enumerate((TI,TE,NP)): c=ws.cell(r,3+j,f"={PL}{L}{row}"); c.number_format='$#,##0;[Red]-$#,##0'; c.border=B
ws.conditional_formatting.add("E14:E25",CellIsRule(operator="lessThan",formula=["0"],font=Font(color="B42318",bold=True)))
head(ws,13,["Top expense categories","Year total"],7)
for i in range(16):
    ws.cell(14+i,7,f"={PL}B{e0+i}"); c=ws.cell(14+i,8,f"={PL}O{e0+i}"); c.number_format='$#,##0'
ws.column_dimensions["H"].width=14
ws.conditional_formatting.add("H14:H29",DataBarRule(start_type="num",start_value=0,end_type="max",color="7FA7D9"))
ch=BarChart(); ch.title="Net profit by month"; ch.add_data(Reference(ws,min_col=5,min_row=13,max_row=25),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=14,max_row=25)); ch.legend=None; ch.height=7.5; ch.width=16
ws.add_chart(ch,"J5")
# Quarterly tax
ws=sheet("Quarterly Tax",[14,16,16,16,18,16,14],"Quarterly Estimated Tax","How much to set aside each quarter")
head(ws,5,["Quarter","Months","Net profit","Set aside","Typical US due date","Paid","Status"])
Q=[("Q1","Jan-Mar","C:E","Apr 15"),("Q2","Apr-May","F:G","Jun 15"),("Q3","Jun-Aug","H:J","Sep 15"),("Q4","Sep-Dec","K:N","Jan 15 (next year)")]
for i,(q,mo,rng,due) in enumerate(Q):
    r=6+i; a,b=rng.split(":"); ws.cell(r,2,q); ws.cell(r,3,mo)
    ws.cell(r,4,f"=SUM({PL}{a}{NP}:{b}{NP})").number_format=MONEY
    ws.cell(r,5,f"=MAX(0,D{r})*Setup!$C$7").number_format=MONEY; ws.cell(r,6,due)
    c=ws.cell(r,7,[1200,1450,0,0][i]); c.fill=fill(NL); c.number_format=MONEY
    ws.cell(r,8,f'=IF(G{r}>=E{r},"Covered",IF(G{r}>0,"Partly paid","To pay"))')
ws.conditional_formatting.add("H6:H9",CellIsRule(operator="equal",formula=['"Covered"'],font=Font(color="067647",bold=True)))
ws.conditional_formatting.add("H6:H9",CellIsRule(operator="equal",formula=['"To pay"'],font=Font(color="B42318",bold=True)))
ws["B11"]="Total"; ws["B11"].font=F(bold=True)
for col in "DEG": ws[f"{col}11"]=f"=SUM({col}6:{col}9)"; ws[f"{col}11"].number_format=MONEY; ws[f"{col}11"].font=F(bold=True)
ws["B13"]="Due dates shown are the usual US federal estimated-tax dates. Confirm yours with your accountant."; ws["B13"].font=F(size=9,color=MUT)
wb.active=wb.sheetnames.index("Dashboard"); wb.save("../out/Small_Business_Bookkeeping_2027.xlsx"); print("ok",TI,TE,NP)
